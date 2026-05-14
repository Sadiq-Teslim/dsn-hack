from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = DOCS / "assets"
EVAL_REPORT = DOCS / "evaluation_report.json"
SUBSET_METRICS = ROOT / "data" / "amazon_subset" / "subset_metrics.json"

BLACK = RGBColor(0, 0, 0)
HEADER_FILL = "E6E6E6"


def load_json(path: Path) -> dict:
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def set_cell_width(cell, width_inches: float) -> None:
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=60, start=95, bottom=60, end=95) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def configure(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.different_first_page_header_footer = True

    styles = document.styles
    for name, size, bold in [
        ("Normal", 9.0, False),
        ("Title", 22, True),
        ("Subtitle", 10.5, False),
        ("Heading 1", 13.5, True),
        ("Heading 2", 11.5, True),
        ("Heading 3", 10.0, True),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = BLACK
        style.paragraph_format.space_after = Pt(2.5)
        style.paragraph_format.line_spacing = 1.0


def add_header_footer(document: Document, label: str) -> None:
    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.text = label
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = BLACK
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Team Ace | DSN x BCT LLM Agent Challenge")
    run.font.size = Pt(8.5)
    run.font.color.rgb = BLACK


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_width(cell, widths[index])
        set_cell_shading(cell, HEADER_FILL)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell.paragraphs[0].runs[0].font.bold = True
    for row_values in rows:
        row = table.add_row()
        for index, value in enumerate(row_values):
            cell = row.cells[index]
            cell.text = value
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2.0)
        paragraph.add_run(item)


def add_numbered(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(2.0)
        paragraph.add_run(item)


def add_image(document: Document, path: Path, caption: str, width: float = 6.4) -> None:
    if not path.exists():
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    caption_para = document.add_paragraph()
    caption_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption_para.add_run(caption)
    run.font.size = Pt(8.2)
    run.font.color.rgb = BLACK


def cover(document: Document, title: str, subtitle: str, task_label: str) -> None:
    for _ in range(2):
        document.add_paragraph()
    paragraph = document.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run(title)

    paragraph = document.add_paragraph(style="Subtitle")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run(subtitle)

    add_image(document, ASSETS / "cover_visual.png", "", width=5.45)

    for index, line in enumerate(
        [
            task_label,
            "Team Ace",
            "Teslim Sadiq | Yasir Oyebo | Abiodun Mark",
            "Contact: sadiqadetola08@gmail.com",
            "Data & AI Summit Hackathon 3.0",
            "DSN x BCT LLM Agent Challenge | May 2026",
        ]
    ):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(1)
        run = paragraph.add_run(line)
        run.bold = index in (0, 1)
        run.font.size = Pt(10.5 if index < 2 else 9)
        run.font.color.rgb = BLACK


def enforce_black(document: Document) -> None:
    parts = [document]
    for section in document.sections:
        parts.extend([section.header, section.footer])
    for part in parts:
        for paragraph in part.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = BLACK
        for table in part.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.color.rgb = BLACK


def dataset_rows(subset: dict) -> list[list[str]]:
    return [
        ["Source", "Amazon Reviews 2023", "Official McAuley Lab review data with bounded local subset"],
        ["Categories", ", ".join(subset.get("categories", [])), "Beauty, grocery, movies, and games behavior"],
        ["Products", str(subset.get("products", "")), "Candidate items available to rating and ranking models"],
        ["Reviews", str(subset.get("reviews", "")), "Real review text and star ratings after filtering"],
        ["Users", str(subset.get("users", "")), "Users with chronological interactions"],
        ["Holdout", f"{subset.get('train_interactions', '')} train / {subset.get('test_interactions', '')} test", "Most recent interaction held out where possible"],
    ]


def build_task_a() -> Path:
    eval_report = load_json(EVAL_REPORT)
    subset = load_json(SUBSET_METRICS)
    path = DOCS / "Task_A_User_Modeling_Team_Ace.docx"
    document = Document()
    configure(document)
    add_header_footer(document, "Task A: User Modeling | Team Ace")
    cover(
        document,
        "Task A: Dynamic User Modeling",
        "Simulating realistic ratings and grounded reviews with auditable LLM agents",
        "Standalone Solution Paper",
    )
    document.add_page_break()

    document.add_heading("Executive Summary", level=1)
    document.add_paragraph(
        "Task A is about behavioral fidelity: given a user persona and an unseen product, the system "
        "must predict the rating the user would give and write a review that sounds consistent with "
        "their preferences, tone, history, and context. Team Ace solves this by treating the user as a "
        "dynamic behavioral persona rather than a static profile label."
    )
    document.add_paragraph(
        "The core design separates decision-making from language generation. A local, interpretable "
        "rating model predicts the star score first. Retrieval then provides grounding evidence from "
        "the user's history. Groq is only used after the rating decision to produce fluent review text. "
        "If Groq is unavailable, a deterministic fallback still returns a complete review response."
    )
    add_table(
        document,
        ["Requirement", "Team Ace Implementation", "Judge Evidence"],
        [
            ["Star rating", "Local rating ensemble using user bias, product signal, category affinity, preference match, and budget pressure", "rating, confidence, reasoning"],
            ["Written review", "Retrieval-augmented Groq prompt with deterministic fallback", "review_text and fallback_used"],
            ["Behavioral fidelity", "Persona memory includes history, tone, likes, dislikes, Nigerian context, and language preference", "evidence array plus reasoning"],
            ["Reproducibility", "Standalone FastAPI app, Docker/Render config, fixture data, and tests", "app.task_a_main:app"],
        ],
        [1.35, 3.3, 1.75],
    )

    document.add_heading("1. Problem and Opportunity", level=1)
    document.add_paragraph(
        "Most review simulators either memorize product averages or ask an LLM to invent a plausible "
        "review. Both approaches miss the point of the challenge. A user's rating is shaped by their "
        "own strictness, their category interests, price sensitivity, previous experiences, and the "
        "moment in which the product is being evaluated. For Nigerian users, this context can include "
        "student budgets, family shopping patterns, heat or harmattan conditions, local media taste, "
        "and the natural mix of English and Nigerian language styles."
    )
    document.add_paragraph(
        "The opportunity is to make review generation measurable and inspectable. The system should not "
        "only produce nice text; it should expose why it believes this user would give this rating."
    )
    document.add_heading("2. Why Team Ace Is Different", level=1)
    add_table(
        document,
        ["Aspect", "Typical Approach", "Team Ace Approach", "Advantage"],
        [
            ["Rating", "LLM guesses the stars", "Local interpretable scoring", "Auditable and testable"],
            ["Review text", "Ungrounded free generation", "Retrieved evidence plus persona prompt", "More behaviorally faithful"],
            ["User model", "Static profile fields", "Dynamic traces from history and context", "Captures changing behavior"],
            ["Nigerian fit", "Generic global shopper", "Budget, family, location, language, and routine signals", "More realistic local output"],
        ],
        [1.1, 1.65, 2.35, 1.4],
    )
    document.add_page_break()

    document.add_heading("3. Architecture and Agentic Workflow", level=1)
    add_image(
        document,
        ASSETS / "architecture.png",
        "Figure 1. Task A uses the shared agent architecture but activates the rating and review path.",
        width=6.35,
    )
    add_numbered(
        document,
        [
            "Read the persona and summarize rating bias, tone, category interests, likes, dislikes, budget level, location, and language preference.",
            "Compare the unseen product against product metadata, category, price, and the user's prior interactions.",
            "Retrieve similar historical reviews as grounding evidence for the generated review.",
            "Predict rating with the local ensemble before any LLM text is requested.",
            "Ask Groq to write the review using the validated score, evidence, persona, and product details.",
            "Return rating, review_text, confidence, reasoning, evidence, llm_provider, and fallback_used.",
        ],
    )
    add_table(
        document,
        ["Layer", "Role in Task A"],
        [
            ["Persona modeling", "Builds the user's behavioral memory from stated profile and review history."],
            ["Retrieval", "Selects similar interactions to ground tone and evidence."],
            ["Rating engine", "Combines local numerical features into a bounded 1-5 star prediction."],
            ["Groq generation", "Turns the validated decision into fluent review text."],
            ["Fallback layer", "Keeps output reproducible when hosted generation is unavailable."],
        ],
        [1.55, 4.95],
    )

    document.add_heading("4. Dataset Strategy and Experiments", level=1)
    document.add_paragraph(
        "The repository uses a zero-download fixture set for immediate judge demos and a bounded real "
        "Amazon Reviews 2023 subset for measurable evaluation. The real subset covers All Beauty, "
        "Grocery and Gourmet Food, Movies and TV, and Video Games. The split holds out each user's "
        "latest interaction where enough history exists."
    )
    add_table(document, ["Dataset Field", "Value", "Meaning"], dataset_rows(subset), [1.3, 2.25, 2.95])
    add_table(
        document,
        ["Experiment", "Purpose", "Expected Signal"],
        [
            ["Global mean baseline", "Tests whether personalization improves beyond average rating", "Lower RMSE"],
            ["Item/category signal", "Checks product and category contribution", "Better fit for unseen items"],
            ["No retrieval evidence", "Tests review grounding", "Less specific and less faithful language"],
            ["No Groq fallback", "Separates scoring reliability from language fluency", "Runnable but less polished text"],
        ],
        [1.55, 2.6, 2.35],
    )
    document.add_page_break()

    document.add_heading("5. Results and Evaluation", level=1)
    add_table(
        document,
        ["Metric", "Team Ace", "Baseline", "Interpretation"],
        [
            [
                "Task A RMSE",
                str(eval_report.get("task_a", {}).get("personalized_rmse", "0.7357")),
                str(eval_report.get("task_a", {}).get("global_mean_rmse", "1.3133")),
                "Personalized behavioral features reduce rating error.",
            ],
            ["Review quality", "Grounded RAG output", "Generic text", "Review style is conditioned on history, persona, and evidence."],
            ["Fallback", "Supported", "Often absent", "Judges can run the endpoint without paid keys."],
        ],
        [1.25, 1.25, 1.25, 2.75],
    )
    document.add_paragraph(
        "The RMSE result matters because it shows that the review simulator is not merely a text demo. "
        "The system makes a behavioral prediction before writing. The generated review is then a "
        "natural-language expression of the model's decision, not the source of the decision itself."
    )
    document.add_heading("6. Nigerian Contextualization", level=1)
    add_table(
        document,
        ["Persona", "Signal", "Effect on Review"],
        [
            ["Lagos student", "Low budget, school routine, transport cost", "More value-sensitive wording and practical tradeoffs"],
            ["Abuja professional", "Work routine, heat, polished daily use", "Reliability and presentation become stronger factors"],
            ["Port Harcourt family shopper", "Visitors, children, home use", "Family-safe, shareable, gentle, and value-pack reasoning"],
        ],
        [1.55, 2.2, 2.75],
    )
    document.add_paragraph(
        "Yarn Mode and YarnGPT audio are output layers for accessibility and local resonance. They do "
        "not change the underlying rating score; they make the explanation easier to present in a "
        "Nigerian demo context."
    )

    document.add_heading("7. Demo, Reproducibility, and Next Steps", level=1)
    add_bullets(
        document,
        [
            "Standalone deployment: python -m uvicorn app.task_a_main:app --host 0.0.0.0 --port $PORT.",
            "Local endpoint: POST /api/v1/generate-review.",
            "Task A UI opens directly at the service root.",
            "Docker Compose exposes Task A on port 8001.",
            "Tests verify the split app serves Task A without exposing the Task B endpoint.",
        ],
    )
    document.add_paragraph(
        "Future work would scale the subset, add embedding retrieval for review evidence, run stronger "
        "BERTScore evaluation, and learn calibration parameters from larger user histories."
    )
    document.add_heading("Team and Conclusion", level=1)
    document.add_paragraph(
        "Team Ace demonstrates that Task A can be fluent, culturally aware, and measurable. The agent "
        "predicts behavior first, then speaks in a grounded user voice."
    )
    enforce_black(document)
    document.save(path)
    return path


def build_task_b() -> Path:
    eval_report = load_json(EVAL_REPORT)
    subset = load_json(SUBSET_METRICS)
    path = DOCS / "Task_B_Recommendation_Team_Ace.docx"
    document = Document()
    configure(document)
    add_header_footer(document, "Task B: Recommendation | Team Ace")
    cover(
        document,
        "Task B: Contextual Recommendation",
        "Ranking cross-domain items with dynamic personas, cold-start logic, and Nigerian context",
        "Standalone Solution Paper",
    )
    document.add_page_break()

    document.add_heading("Executive Summary", level=1)
    document.add_paragraph(
        "Task B asks the agent to predict what a user will choose next. Team Ace solves this with a "
        "local recommendation engine that ranks products before any LLM explanation is generated. The "
        "ranker uses persona memory, current context, category affinity, preference match, popularity, "
        "budget fit, dislike penalties, diversity, and cold-start fallbacks."
    )
    document.add_paragraph(
        "The system is designed for the hard cases in the rubric: cold-start, cross-domain, contextual "
        "relevance, and conversational explanation. Groq improves the natural-language summary after "
        "the ranked list is already produced. Yarn Mode can then localize the explanation for Nigerian "
        "Pidgin or Nigerian-flavoured English voice output."
    )
    add_table(
        document,
        ["Requirement", "Team Ace Implementation", "Judge Evidence"],
        [
            ["Personalized ranking", "Scores candidate products with persona and context features", "ranked items, scores, reasons"],
            ["Cold-start", "Uses stated preferences, budget, category priors, and popularity when history is sparse", "demo personas and top_k output"],
            ["Cross-domain", "Ranks across beauty, grocery, movies, and games", "include_categories and catalog candidates"],
            ["Reproducibility", "Local ranking plus deterministic fallback explanations", "app.task_b_main:app"],
        ],
        [1.35, 3.3, 1.75],
    )

    document.add_heading("1. Problem and Opportunity", level=1)
    document.add_paragraph(
        "A recommendation system should do more than repeat popular items. Real users move across "
        "domains in the same week: food for school, a film for the weekend, a game for friends, or a "
        "beauty product for weather and routine. Nigerian context makes this more interesting because "
        "budget, transport, family use, local media taste, and language can all change what feels like "
        "a good recommendation."
    )
    document.add_paragraph(
        "Pure collaborative filtering struggles when the user is new or the context changes. Pure LLM "
        "recommendation can sound persuasive but may not produce stable rankings. Team Ace combines "
        "interpretable ranking with controlled LLM explanation."
    )
    document.add_heading("2. Why Team Ace Is Different", level=1)
    add_table(
        document,
        ["Aspect", "Typical Approach", "Team Ace Approach", "Advantage"],
        [
            ["Ranking", "Popularity or opaque LLM output", "Local scoring before explanation", "Auditable recommendation order"],
            ["Cold-start", "Weak without history", "Uses persona fields and context text", "Useful from first interaction"],
            ["Cross-domain", "Single catalog type", "Beauty, grocery, movies, games", "Better reflects real routines"],
            ["Explanation", "Generic reason text", "Groq explains validated ranked list", "Fluent but controlled"],
        ],
        [1.1, 1.65, 2.35, 1.4],
    )
    document.add_page_break()

    document.add_heading("3. Architecture and Agentic Workflow", level=1)
    add_image(
        document,
        ASSETS / "architecture.png",
        "Figure 1. Task B uses the shared architecture but activates candidate generation and ranking.",
        width=6.35,
    )
    add_numbered(
        document,
        [
            "Read the persona and summarize interests, budget, dislikes, category affinity, cultural context, and history.",
            "Read the current context prompt and extract immediate need signals.",
            "Generate candidates from the catalog across allowed categories.",
            "Score each candidate using popularity, average rating, category fit, text match, preference match, budget fit, and dislike penalties.",
            "Apply ranking and diversity logic to produce the top recommendations.",
            "Ask Groq to summarize the already-ranked list; use deterministic fallback when unavailable.",
        ],
    )
    add_table(
        document,
        ["Layer", "Role in Task B"],
        [
            ["Persona modeling", "Turns user history and stated profile into ranking features."],
            ["Candidate generation", "Builds a cross-domain pool from the local catalog."],
            ["Ranking engine", "Scores items locally and returns deterministic ordering."],
            ["Explanation layer", "Uses Groq to explain the ranked list without changing it."],
            ["Voice layer", "Uses Yarn Mode and YarnGPT for local-language presentation."],
        ],
        [1.55, 4.95],
    )

    document.add_heading("4. Dataset Strategy and Experiments", level=1)
    document.add_paragraph(
        "The recommender uses the same two data tracks as the full submission: curated fixtures for "
        "instant demos and a bounded real Amazon Reviews 2023 subset for measurable ranking tests. The "
        "holdout strategy treats each user's latest interaction as the target when possible."
    )
    add_table(document, ["Dataset Field", "Value", "Meaning"], dataset_rows(subset), [1.3, 2.25, 2.95])
    add_table(
        document,
        ["Experiment", "Purpose", "Expected Signal"],
        [
            ["Popularity baseline", "Measures whether personalization beats generic item popularity", "Higher NDCG@10 and Hit Rate@10"],
            ["Cold-start personas", "Checks if recommendations work with sparse history", "Reasonable rankings from stated context"],
            ["Cross-domain ranking", "Tests ranking across all selected Amazon categories", "Useful mixed-category lists"],
            ["No Groq explanation", "Separates ranking quality from language polish", "Same ranking remains available"],
        ],
        [1.55, 2.6, 2.35],
    )
    document.add_page_break()

    document.add_heading("5. Results and Evaluation", level=1)
    add_image(
        document,
        ASSETS / "ablation_results.png",
        "Figure 2. Personalization improves over the popularity baseline on the checked-in subset.",
        width=5.95,
    )
    add_table(
        document,
        ["Metric", "Team Ace", "Baseline", "Interpretation"],
        [
            [
                "NDCG@10",
                str(eval_report.get("task_b", {}).get("personalized_ndcg_at_10", "0.0481")),
                str(eval_report.get("task_b", {}).get("popularity_ndcg_at_10", "0.0194")),
                "Personalized ranking places held-out choices higher.",
            ],
            [
                "Hit Rate@10",
                str(eval_report.get("task_b", {}).get("personalized_hit_rate_at_10", "0.1163")),
                str(eval_report.get("task_b", {}).get("popularity_hit_rate_at_10", "0.0349")),
                "More held-out items appear in the top 10.",
            ],
            ["Explanation fallback", "Supported", "Often absent", "The ranked list is available without paid keys."],
        ],
        [1.25, 1.25, 1.25, 2.75],
    )
    document.add_paragraph(
        "The absolute ranking scores are modest because the committed subset is small and sparse. The "
        "important signal is the ablation direction: the local personalized ranker improves on a "
        "popularity-only baseline while keeping the ordering inspectable."
    )
    document.add_heading("6. Nigerian Contextualization", level=1)
    add_table(
        document,
        ["Persona", "Need Context", "Recommendation Effect"],
        [
            ["Lagos student", "Busy school week and low budget", "Affordable quick meals, multiplayer games, and local entertainment"],
            ["Abuja professional", "Weekday work routine and relaxed evenings", "Dependable coffee, polished media, practical skincare"],
            ["Port Harcourt family shopper", "Home use, children, visitors", "Family-safe media, guest-friendly drinks, gentle beauty products"],
        ],
        [1.55, 2.2, 2.75],
    )
    document.add_paragraph(
        "The Nigerian layer is not decoration. It changes recommendation reasons around affordability, "
        "shared use, family setting, weather, work routines, and language preference."
    )

    document.add_heading("7. Demo, Reproducibility, and Next Steps", level=1)
    add_bullets(
        document,
        [
            "Standalone deployment: python -m uvicorn app.task_b_main:app --host 0.0.0.0 --port $PORT.",
            "Local endpoint: POST /api/v1/recommend.",
            "Task B UI opens directly at the service root.",
            "Docker Compose exposes Task B on port 8002.",
            "Tests verify the split app serves Task B without exposing the Task A endpoint.",
        ],
    )
    document.add_paragraph(
        "Future work would add denser metadata, embedding candidate retrieval, larger four-domain "
        "evaluation, multi-turn conversational memory, and Nigerian marketplace data."
    )
    document.add_heading("Team and Conclusion", level=1)
    document.add_paragraph(
        "Team Ace demonstrates a recommender that reasons before it speaks. Task B is contextual, "
        "cross-domain, reproducible, and ready to be deployed as its own service."
    )
    enforce_black(document)
    document.save(path)
    return path


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    for path in [build_task_a(), build_task_b()]:
        print(path)


if __name__ == "__main__":
    main()
